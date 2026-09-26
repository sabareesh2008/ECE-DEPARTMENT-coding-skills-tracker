"""
Generate exact template match PowerPoint presentation for CODEMETRIX.
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# Exact Template Colors
BG_DARK = RGBColor(8, 18, 37)        # #081225 Deep Navy
BG_LIGHT = RGBColor(248, 250, 252)   # #F8FAFC Light Slate
CARD_WHITE = RGBColor(255, 255, 255) # Pure White
CARD_NAVY = RGBColor(11, 30, 54)     # #0B1E36 Dark Blue Card
BORDER_LIGHT = RGBColor(226, 232, 240) # #E2E8F0 Border
BORDER_NAVY = RGBColor(30, 41, 59)   # #1E293B
SKY_BLUE = RGBColor(14, 165, 233)    # #0EA5E9 Accent Sky
TEXT_DARK = RGBColor(15, 23, 42)     # #0F172A Slate 900
TEXT_MUTED = RGBColor(100, 116, 139) # #64748B Slate 500
TEXT_WHITE = RGBColor(248, 250, 252) # #F8FAFC

def set_slide_bg(slide, is_dark=False):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG_DARK if is_dark else BG_LIGHT
    bg.line.fill.background()
    return bg

def add_header(slide, kicker, title, subtitle=None):
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(10.5), Inches(1.1))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    
    p0 = tf.paragraphs[0]
    p0.text = kicker.upper()
    p0.font.name = "Segoe UI"
    p0.font.size = Pt(9.5)
    p0.font.bold = True
    p0.font.color.rgb = SKY_BLUE
    p0.space_after = Pt(2)
    
    p1 = tf.add_paragraph()
    p1.text = title
    p1.font.name = "Segoe UI"
    p1.font.size = Pt(22)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_DARK
    
    if subtitle:
        p2 = tf.add_paragraph()
        p2.text = subtitle
        p2.font.name = "Segoe UI"
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = TEXT_MUTED

def add_footer(slide, slide_num, is_dark=False):
    # Left Brand
    tb_l = slide.shapes.add_textbox(Inches(0.8), Inches(6.9), Inches(4), Inches(0.4))
    tf_l = tb_l.text_frame
    tf_l.margin_left = tf_l.margin_top = 0
    p_l = tf_l.paragraphs[0]
    p_l.text = "C O D E M E T R I X"
    p_l.font.name = "Segoe UI"
    p_l.font.size = Pt(9)
    p_l.font.bold = True
    p_l.font.color.rgb = RGBColor(71, 85, 105) if is_dark else RGBColor(148, 163, 184)
    
    # Right Page Num
    tb_r = slide.shapes.add_textbox(Inches(11.5), Inches(6.9), Inches(1.0), Inches(0.4))
    tf_r = tb_r.text_frame
    tf_r.margin_left = tf_r.margin_top = 0
    p_r = tf_r.paragraphs[0]
    p_r.text = f"{slide_num:02d}"
    p_r.font.name = "Segoe UI"
    p_r.font.size = Pt(9)
    p_r.font.bold = True
    p_r.font.color.rgb = RGBColor(71, 85, 105) if is_dark else RGBColor(148, 163, 184)
    p_r.alignment = PP_ALIGN.RIGHT

def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    
    # ==================== SLIDE 1 ====================
    s1 = prs.slides.add_slide(blank)
    set_slide_bg(s1, is_dark=True)
    
    # Top Icon
    tb_top = s1.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(3.0), Inches(0.5))
    tf_top = tb_top.text_frame
    p_top = tf_top.paragraphs[0]
    p_top.text = "<>  CODEMETRIX"
    p_top.font.name = "Segoe UI"
    p_top.font.size = Pt(11)
    p_top.font.bold = True
    p_top.font.color.rgb = SKY_BLUE
    
    # Left Content
    tb1 = s1.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(7.5), Inches(4.5))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    
    p = tf1.paragraphs[0]
    p.text = "CODEMETRIX"
    p.font.name = "Segoe UI"
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE
    p.space_after = Pt(4)
    
    p = tf1.add_paragraph()
    p.text = "From Coding Activity to Actionable Performance Insights"
    p.font.name = "Segoe UI"
    p.font.size = Pt(14)
    p.font.italic = True
    p.font.color.rgb = SKY_BLUE
    p.space_after = Pt(36)
    
    # Meta Items
    rows = [
        ("TEAM NAME", "CodeMetrix"),
        ("TEAM MEMBERS", "Sabareesh K  •  Sabarivasan P  •  Roshan M"),
        ("DEPARTMENT", "II ECE E")
    ]
    for label, val in rows:
        p_row = tf1.add_paragraph()
        r1 = p_row.add_run()
        r1.text = f"{label:<16}  "
        r1.font.name = "Segoe UI"
        r1.font.size = Pt(10)
        r1.font.bold = True
        r1.font.color.rgb = SKY_BLUE
        
        r2 = p_row.add_run()
        r2.text = val
        r2.font.name = "Segoe UI"
        r2.font.size = Pt(11)
        r2.font.color.rgb = TEXT_WHITE
        p_row.space_after = Pt(12)
        
    # Right QR Box
    qr_card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.2), Inches(1.8), Inches(3.3), Inches(4.5))
    qr_card.fill.solid()
    qr_card.fill.fore_color.rgb = CARD_WHITE
    qr_card.line.fill.background()
    
    tf_qr = qr_card.text_frame
    tf_qr.word_wrap = True
    tf_qr.margin_top = Inches(0.4)
    
    p_q1 = tf_qr.paragraphs[0]
    p_q1.text = "▣"
    p_q1.font.name = "Segoe UI"
    p_q1.font.size = Pt(90)
    p_q1.font.color.rgb = TEXT_DARK
    p_q1.alignment = PP_ALIGN.CENTER
    p_q1.space_after = Pt(12)
    
    p_q2 = tf_qr.add_paragraph()
    p_q2.text = "SCAN TO VIEW\nCODEMETRIX"
    p_q2.font.name = "Segoe UI"
    p_q2.font.size = Pt(10)
    p_q2.font.bold = True
    p_q2.font.color.rgb = TEXT_DARK
    p_q2.alignment = PP_ALIGN.CENTER
    
    add_footer(s1, 1, is_dark=True)
    
    # ==================== SLIDE 2 ====================
    s2 = prs.slides.add_slide(blank)
    set_slide_bg(s2)
    add_header(s2, "THE CHALLENGE", "Problem Statement", "Students practice coding on platforms like LeetCode and GitHub, but faculty often lack a centralized system to track and evaluate that progress effectively.")
    
    # 3x2 Grid Cards
    cards_s2 = [
        ("Manual Tracking", "Manual tracking of students' coding progress across individual profiles."),
        ("Hard to Monitor at Scale", "Difficult to monitor hundreds of students and keep records continually updated."),
        ("Inactive Students Missed", "Faculty cannot easily identify inactive students early for timely guidance."),
        ("Scattered Data & Copy-Paste", "Student performance is scattered without code authenticity and paste verification."),
        ("No Easy Comparison", "Difficult to compare section-wise performance, daily streaks, and authentic logic."),
        ("Database Storage Limits", "Saving raw code submissions quickly exhausts free cloud limits without an auto-purge.")
    ]
    for i, (title, desc) in enumerate(cards_s2):
        row, col = i // 3, i % 3
        x = 0.8 + col * 3.95
        y = 2.0 + row * 2.3
        
        c = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(3.8), Inches(2.1))
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_WHITE
        c.line.color.rgb = BORDER_LIGHT
        
        tf = c.text_frame
        tf.word_wrap = True
        tf.margin_top = Inches(0.25)
        tf.margin_left = Inches(0.25)
        tf.margin_right = Inches(0.25)
        
        pt = tf.paragraphs[0]
        pt.text = title
        pt.font.name = "Segoe UI"
        pt.font.size = Pt(12)
        pt.font.bold = True
        pt.font.color.rgb = TEXT_DARK
        pt.space_after = Pt(6)
        
        pd = tf.add_paragraph()
        pd.text = desc
        pd.font.name = "Segoe UI"
        pd.font.size = Pt(9.5)
        pd.font.color.rgb = TEXT_MUTED
        
    add_footer(s2, 2)
    
    # ==================== SLIDE 3 ====================
    s3 = prs.slides.add_slide(blank)
    set_slide_bg(s3)
    add_header(s3, "THE APPROACH", "Proposed Solution", "CodeMetrix is a centralized dashboard that automatically collects and analyzes students' LeetCode and GitHub activity, and presents it through student, section, and faculty dashboards.")
    
    # Top 4 Navy boxes
    top_boxes = [
        "LeetCode & GitHub",
        "Automated Data Collection",
        "Data Processing & Forensics",
        "Supabase Database"
    ]
    for i, name in enumerate(top_boxes):
        x = 0.8 + i * 2.95
        b = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.0), Inches(2.8), Inches(1.3))
        b.fill.solid()
        b.fill.fore_color.rgb = CARD_NAVY
        b.line.color.rgb = BORDER_NAVY
        tf = b.text_frame
        tf.word_wrap = True
        tf.margin_top = Inches(0.35)
        p = tf.paragraphs[0]
        p.text = name
        p.font.name = "Segoe UI"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE
        p.alignment = PP_ALIGN.CENTER
        
    # Middle Pill
    pill = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.5), Inches(3.7), Inches(4.333), Inches(0.7))
    pill.fill.solid()
    pill.fill.fore_color.rgb = SKY_BLUE
    pill.line.fill.background()
    tf_p = pill.text_frame
    p_p = tf_p.paragraphs[0]
    p_p.text = "CODEMETRIX"
    p_p.font.name = "Segoe UI"
    p_p.font.size = Pt(14)
    p_p.font.bold = True
    p_p.font.color.rgb = TEXT_DARK
    p_p.alignment = PP_ALIGN.CENTER
    
    # Bottom 3 White boxes
    bottom_boxes = ["Student Dashboard", "Faculty Analytics", "Admin Dashboard"]
    for i, name in enumerate(bottom_boxes):
        x = 0.8 + i * 3.95
        b = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(4.8), Inches(3.8), Inches(1.5))
        b.fill.solid()
        b.fill.fore_color.rgb = CARD_WHITE
        b.line.color.rgb = BORDER_LIGHT
        tf = b.text_frame
        tf.word_wrap = True
        tf.margin_top = Inches(0.5)
        p = tf.paragraphs[0]
        p.text = name
        p.font.name = "Segoe UI"
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = TEXT_DARK
        p.alignment = PP_ALIGN.CENTER
        
    add_footer(s3, 3)
    
    # ==================== SLIDE 4 ====================
    s4 = prs.slides.add_slide(blank)
    set_slide_bg(s4)
    add_header(s4, "UNDER THE HOOD", "Technology Stack")
    
    cols_s4 = [
        ("FRONTEND", [
            ("HTML", "Builds the structure of the web dashboard"),
            ("CSS", "Creates the responsive and modern UI"),
            ("JavaScript & JSZip", "Handles interactions, filtering, search & 1-click ZIP downloads")
        ]),
        ("AUTOMATION & PROCESSING", [
            ("Python", "Automates LeetCode/GitHub data collection and processing"),
            ("GitHub Actions", "Runs automated trackers periodically without manual execution"),
            ("Monaco DOM Hook", "Browser extension capturing typing rhythm and paste counts")
        ]),
        ("DATABASE & AUTHENTICATION", [
            ("SQL & Triggers", "Auto-replace deduplication constraints & 50-problem purge"),
            ("Supabase", "Stores student profiles, coding data and application data"),
            ("Row Level Security", "Guarantees secure access across faculty and student portals")
        ])
    ]
    for col_idx, (col_title, items) in enumerate(cols_s4):
        x = 0.8 + col_idx * 3.95
        card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.8), Inches(3.8), Inches(4.8))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_WHITE
        card.line.color.rgb = BORDER_LIGHT
        
        # Header bar
        hbar = s4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(1.8), Inches(3.8), Inches(0.6))
        hbar.fill.solid()
        hbar.fill.fore_color.rgb = CARD_NAVY
        hbar.line.fill.background()
        tf_h = hbar.text_frame
        p_h = tf_h.paragraphs[0]
        p_h.text = col_title
        p_h.font.name = "Segoe UI"
        p_h.font.size = Pt(10)
        p_h.font.bold = True
        p_h.font.color.rgb = TEXT_WHITE
        p_h.alignment = PP_ALIGN.CENTER
        
        # Content
        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_top = Inches(0.8)
        tf.margin_left = Inches(0.25)
        tf.margin_right = Inches(0.25)
        
        for name, desc in items:
            p_t = tf.add_paragraph()
            p_t.text = name
            p_t.font.name = "Segoe UI"
            p_t.font.size = Pt(11)
            p_t.font.bold = True
            p_t.font.color.rgb = TEXT_DARK
            
            p_d = tf.add_paragraph()
            p_d.text = desc
            p_d.font.name = "Segoe UI"
            p_d.font.size = Pt(9)
            p_d.font.color.rgb = TEXT_MUTED
            p_d.space_after = Pt(10)
            
    add_footer(s4, 4)
    
    # ==================== SLIDE 5 ====================
    s5 = prs.slides.add_slide(blank)
    set_slide_bg(s5)
    add_header(s5, "HOW IT WORKS", "System Architecture")
    
    # Tier 1
    t1_1 = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.7), Inches(0.9))
    t1_1.fill.solid()
    t1_1.fill.fore_color.rgb = CARD_NAVY
    t1_1.line.color.rgb = BORDER_NAVY
    tf = t1_1.text_frame
    tf.margin_top = Inches(0.12)
    tf.margin_left = Inches(0.2)
    p = tf.paragraphs[0]
    p.text = "LeetCode"
    p.font.bold = True
    p.font.size = Pt(11)
    p.font.color.rgb = SKY_BLUE
    p = tf.add_paragraph()
    p.text = "• Problems solved (E / M / H)   • Daily activity & submissions"
    p.font.size = Pt(9.5)
    p.font.color.rgb = TEXT_WHITE
    
    t1_2 = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(0.9))
    t1_2.fill.solid()
    t1_2.fill.fore_color.rgb = CARD_NAVY
    t1_2.line.color.rgb = BORDER_NAVY
    tf = t1_2.text_frame
    tf.margin_top = Inches(0.12)
    tf.margin_left = Inches(0.2)
    p = tf.paragraphs[0]
    p.text = "GitHub"
    p.font.bold = True
    p.font.size = Pt(11)
    p.font.color.rgb = SKY_BLUE
    p = tf.add_paragraph()
    p.text = "• Repositories & commits   • Contribution activity"
    p.font.size = Pt(9.5)
    p.font.color.rgb = TEXT_WHITE
    
    # Tier 2
    t2 = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.9), Inches(11.7), Inches(0.9))
    t2.fill.solid()
    t2.fill.fore_color.rgb = CARD_WHITE
    t2.line.color.rgb = BORDER_LIGHT
    tf = t2.text_frame
    tf.margin_top = Inches(0.12)
    tf.margin_left = Inches(0.2)
    p = tf.paragraphs[0]
    p.text = "Python — Processing Engine"
    p.font.bold = True
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_DARK
    p = tf.add_paragraph()
    p.text = "• Fetch data   • Process and clean data   • Validate deduplication   • Enforce 50-limit window"
    p.font.size = Pt(9.5)
    p.font.color.rgb = TEXT_MUTED
    
    # Tier 3
    t3_1 = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.0), Inches(5.7), Inches(0.9))
    t3_1.fill.solid()
    t3_1.fill.fore_color.rgb = CARD_WHITE
    t3_1.line.color.rgb = BORDER_LIGHT
    tf = t3_1.text_frame
    tf.margin_top = Inches(0.12)
    tf.margin_left = Inches(0.2)
    p = tf.paragraphs[0]
    p.text = "GitHub Actions"
    p.font.bold = True
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_DARK
    p = tf.add_paragraph()
    p.text = "• Scheduled execution   • Runs trackers automatically"
    p.font.size = Pt(9.5)
    p.font.color.rgb = TEXT_MUTED
    
    t3_2 = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(4.0), Inches(5.7), Inches(0.9))
    t3_2.fill.solid()
    t3_2.fill.fore_color.rgb = CARD_WHITE
    t3_2.line.color.rgb = BORDER_LIGHT
    tf = t3_2.text_frame
    tf.margin_top = Inches(0.12)
    tf.margin_left = Inches(0.2)
    p = tf.paragraphs[0]
    p.text = "Supabase (PostgreSQL)"
    p.font.bold = True
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_DARK
    p = tf.add_paragraph()
    p.text = "• Central data store   • Auto-prune triggers   • Powers all dashboards"
    p.font.size = Pt(9.5)
    p.font.color.rgb = TEXT_MUTED
    
    # Tier 4
    dashboards = [
        ("Student Dashboard", "• LeetCode & GitHub activity\n• 1-Click ZIP archive download"),
        ("Admin Dashboard", "• Manage profiles & faculty\n• Anti-paste forensics inspection"),
        ("Faculty Dashboard", "• Section-wise analysis\n• In-modal code inspection & Excel export")
    ]
    for i, (title, desc) in enumerate(dashboards):
        x = 0.8 + i * 3.95
        d = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(5.1), Inches(3.8), Inches(1.5))
        d.fill.solid()
        d.fill.fore_color.rgb = CARD_NAVY
        d.line.color.rgb = BORDER_NAVY
        tf = d.text_frame
        tf.margin_top = Inches(0.2)
        tf.margin_left = Inches(0.2)
        p = tf.paragraphs[0]
        p.text = title
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = SKY_BLUE
        p = tf.add_paragraph()
        p.text = desc
        p.font.size = Pt(9)
        p.font.color.rgb = TEXT_WHITE
        
    add_footer(s5, 5)
    
    # ==================== SLIDE 6 ====================
    s6 = prs.slides.add_slide(blank)
    set_slide_bg(s6)
    add_header(s6, "WHAT IT DELIVERS", "Key Features")
    
    cards_s6 = [
        ("Automated LeetCode & GitHub Tracking", "Daily background scraping captures solved counts, streaks, and difficulty breakdowns automatically."),
        ("Section-wise Analysis", "Instant filtering across sections with ranked leaderboards and detailed student modals."),
        ("Inactive Students Detection", "Early identification of students with broken practice streaks for prompt academic mentoring."),
        ("Suspicious Solving Analysis", "Extension tracks typing vs paste deltas to flag AI-pasted code while keeping short one-liners clean."),
        ("50-Problem Download & Auto-Purge", "1-click ZIP and Excel archive of solved files with automatic database reset to maintain free cloud limits."),
        ("Duplicate Prevention", "Re-submitted problems automatically replace older records, guaranteeing zero duplicate entries.")
    ]
    for i, (title, desc) in enumerate(cards_s6):
        row, col = i // 3, i % 3
        x = 0.8 + col * 3.95
        y = 1.8 + row * 2.4
        
        c = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(3.8), Inches(2.2))
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_WHITE
        c.line.color.rgb = BORDER_LIGHT
        
        tf = c.text_frame
        tf.word_wrap = True
        tf.margin_top = Inches(0.25)
        tf.margin_left = Inches(0.25)
        tf.margin_right = Inches(0.25)
        
        pt = tf.paragraphs[0]
        pt.text = title
        pt.font.name = "Segoe UI"
        pt.font.size = Pt(12)
        pt.font.bold = True
        pt.font.color.rgb = TEXT_DARK
        pt.space_after = Pt(6)
        
        pd = tf.add_paragraph()
        pd.text = desc
        pd.font.name = "Segoe UI"
        pd.font.size = Pt(9.5)
        pd.font.color.rgb = TEXT_MUTED
        
    add_footer(s6, 6)
    
    # ==================== SLIDE 7 ====================
    s7 = prs.slides.add_slide(blank)
    set_slide_bg(s7)
    add_header(s7, "WHAT SETS US APART", "Innovation")
    
    cards_s7 = [
        ("Unified Coding Intelligence", "Combines LeetCode + GitHub activity into one intelligence layer."),
        ("Real-Time Monitoring", "Automated, real-time activity monitoring across platforms."),
        ("Pattern-Based Integrity Checks", "Physical keystroke counting and smart one-liner tolerance (<=220 chars) eliminate false alarms."),
        ("Self-Pruning Cloud Architecture", "Maintains perpetual free-tier operation via automated rolling 50-problem pruning and offline archives."),
        ("Data-Driven Reporting", "Historical and section-wise data-driven reporting formatted for reviews."),
        ("Automated Delivery", "Automatic report sending and email updates for class advisors.")
    ]
    for i, (title, desc) in enumerate(cards_s7):
        row, col = i // 2, i % 2
        x = 0.8 + col * 5.95
        y = 1.8 + row * 1.6
        
        c = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(5.8), Inches(1.4))
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_WHITE
        c.line.color.rgb = BORDER_LIGHT
        
        tf = c.text_frame
        tf.word_wrap = True
        tf.margin_top = Inches(0.18)
        tf.margin_left = Inches(0.25)
        tf.margin_right = Inches(0.25)
        
        pt = tf.paragraphs[0]
        pt.text = title
        pt.font.name = "Segoe UI"
        pt.font.size = Pt(11.5)
        pt.font.bold = True
        pt.font.color.rgb = TEXT_DARK
        pt.space_after = Pt(4)
        
        pd = tf.add_paragraph()
        pd.text = desc
        pd.font.name = "Segoe UI"
        pd.font.size = Pt(9.5)
        pd.font.color.rgb = TEXT_MUTED
        
    add_footer(s7, 7)
    
    # ==================== SLIDE 8 ====================
    s8 = prs.slides.add_slide(blank)
    set_slide_bg(s8)
    add_header(s8, "WHY IT MATTERS", "Impact")
    
    rows_s8 = [
        ("Consistent Practice", "Daily coding habit fostered across all students"),
        ("Reduced Workload", "Automated monitoring saves 15+ faculty hours weekly"),
        ("Early Intervention", "Inactive students identified quickly for counseling"),
        ("Data-Driven Decisions", "Meaningful analytics for section-level performance reviews"),
        ("Academic Integrity", "Unusual solving patterns and copy-paste dumps flagged"),
        ("Storage Sustainability", "50-problem auto-pruning enables perpetual zero-cost operation"),
        ("Career Readiness", "Verified LeetCode + GitHub portfolios for placement drives")
    ]
    
    # Table Box
    t_shape = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.733), Inches(4.8))
    t_shape.fill.solid()
    t_shape.fill.fore_color.rgb = CARD_WHITE
    t_shape.line.color.rgb = BORDER_LIGHT
    
    # Header bar
    hbar = s8.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.733), Inches(0.6))
    hbar.fill.solid()
    hbar.fill.fore_color.rgb = CARD_NAVY
    hbar.line.fill.background()
    tf_h = hbar.text_frame
    p_h = tf_h.paragraphs[0]
    p_h.text = f"{'IMPACT':<35}RESULT"
    p_h.font.name = "Segoe UI"
    p_h.font.size = Pt(10)
    p_h.font.bold = True
    p_h.font.color.rgb = TEXT_WHITE
    
    tf = t_shape.text_frame
    tf.word_wrap = True
    tf.margin_top = Inches(0.7)
    tf.margin_left = Inches(0.3)
    
    for imp, res in rows_s8:
        p_row = tf.add_paragraph()
        r1 = p_row.add_run()
        r1.text = f"{imp:<30} "
        r1.font.name = "Segoe UI"
        r1.font.size = Pt(10)
        r1.font.bold = True
        r1.font.color.rgb = TEXT_DARK
        
        r2 = p_row.add_run()
        r2.text = res
        r2.font.name = "Segoe UI"
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = TEXT_MUTED
        p_row.space_after = Pt(6)
        
    add_footer(s8, 8)
    
    # ==================== SLIDE 9 ====================
    s9 = prs.slides.add_slide(blank)
    set_slide_bg(s9)
    add_header(s9, "WHAT'S NEXT", "Future Scope")
    
    future_items = [
        "AI-Based Performance & Complexity Optimization Hints",
        "Advanced Anomaly & Code Quality Detection",
        "Complete Online Department Coding Test Platform",
        "Mobile App & Automated Telegram/WhatsApp Notifications",
        "Integration with Multiple Coding Platforms (Codeforces, HackerRank, GFG)"
    ]
    for i, item in enumerate(future_items):
        y = 1.8 + i * 0.85
        bar = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(y), Inches(11.733), Inches(0.7))
        bar.fill.solid()
        bar.fill.fore_color.rgb = CARD_WHITE
        bar.line.color.rgb = BORDER_LIGHT
        tf = bar.text_frame
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.18)
        p = tf.paragraphs[0]
        p.text = f"•  {item}"
        p.font.name = "Segoe UI"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = TEXT_DARK
        
    # Bottom Vision Callout
    vision = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.1), Inches(11.733), Inches(0.6))
    vision.fill.solid()
    vision.fill.fore_color.rgb = CARD_NAVY
    vision.line.fill.background()
    tf_v = vision.text_frame
    p_v = tf_v.paragraphs[0]
    p_v.text = "The future vision of CodeMetrix is to evolve from a coding activity tracker into an intelligent, personalized coding education platform."
    p_v.font.name = "Segoe UI"
    p_v.font.size = Pt(10)
    p_v.font.italic = True
    p_v.font.color.rgb = TEXT_WHITE
    p_v.alignment = PP_ALIGN.CENTER
    
    add_footer(s9, 9)
    
    # ==================== SLIDE 10 ====================
    s10 = prs.slides.add_slide(blank)
    set_slide_bg(s10, is_dark=True)
    
    # Icon
    icon_box = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.166), Inches(2.2), Inches(1.0), Inches(1.0))
    icon_box.fill.solid()
    icon_box.fill.fore_color.rgb = CARD_NAVY
    icon_box.line.color.rgb = SKY_BLUE
    tf_ic = icon_box.text_frame
    p_ic = tf_ic.paragraphs[0]
    p_ic.text = "<>"
    p_ic.font.name = "Segoe UI"
    p_ic.font.size = Pt(20)
    p_ic.font.bold = True
    p_ic.font.color.rgb = SKY_BLUE
    p_ic.alignment = PP_ALIGN.CENTER
    
    # Text
    tb10 = s10.shapes.add_textbox(Inches(1.5), Inches(3.4), Inches(10.333), Inches(2.0))
    tf10 = tb10.text_frame
    p = tf10.paragraphs[0]
    p.text = "THANK YOU"
    p.font.name = "Segoe UI"
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE
    p.alignment = PP_ALIGN.CENTER
    p.space_after = Pt(10)
    
    p = tf10.add_paragraph()
    p.text = "CODEMETRIX  •  From Coding Activity to Actionable Performance Insights"
    p.font.name = "Segoe UI"
    p.font.size = Pt(13)
    p.font.color.rgb = SKY_BLUE
    p.alignment = PP_ALIGN.CENTER
    
    add_footer(s10, 10, is_dark=True)
    
    out_pptx = os.path.abspath("CodeMetrix_Presentation.pptx")
    prs.save(out_pptx)
    prs.save(os.path.abspath("ECE_CodeMetrix_Presentation.pptx"))
    print(f"PPTX saved successfully to {out_pptx}")

if __name__ == "__main__":
    build()
